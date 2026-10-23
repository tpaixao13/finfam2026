"""Exportação de despesas para .xlsx e .pdf (bytes)."""
import io

from app.services.finance import brl

HEADERS = ['Pessoa', 'Dia', 'Mês', 'Ano', 'Descrição', 'Categoria',
           'Forma de Pagamento', 'Banco', 'Valor (R$)', 'Parcela', 'Recorrente']


def expense_rows(expenses) -> list[list]:
    rows = []
    for e in expenses:
        parcela = ''
        if e.installment_group_id and e.installment_group:
            parcela = f'{e.installment_number}/{e.installment_group.num_installments}'
        elif e.recurring_group_id and e.recurring_group:
            parcela = f'Recorrente {e.recurring_number}/{e.recurring_group.num_recurrences}'
        rows.append([
            e.user.name if e.user else '', e.day, e.display_month, e.display_year,
            e.description, e.category, e.payment_method, e.bank or '',
            float(e.amount), parcela, 'Sim' if e.recurring_group_id else 'Não',
        ])
    return rows


def to_xlsx(rows: list[list]) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Despesas'
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0D6EFD')
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        display = list(row)
        display[8] = brl(row[8])
        ws.append(display)
    for i, width in enumerate([14, 6, 6, 7, 42, 16, 20, 16, 16, 13, 11], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _latin1(value) -> str:
    """As fontes padrão do PDF são latin-1: troca caracteres fora dele."""
    return str(value).replace('•', '*').encode('latin-1', 'replace').decode('latin-1')


def to_pdf(rows: list[list], period_label: str) -> bytes:
    from fpdf import FPDF

    class _PDF(FPDF):
        def header(self):
            self.set_font('Helvetica', 'B', 13)
            self.cell(0, 10, _latin1(f'Despesas - {period_label}'), align='C', new_x='LMARGIN', new_y='NEXT')
            self.ln(3)

        def footer(self):
            self.set_y(-12)
            self.set_font('Helvetica', 'I', 8)
            self.cell(0, 10, f'Pagina {self.page_no()}', align='C')

    pdf = _PDF(orientation='L', unit='mm', format='A4')
    pdf.set_margins(10, 10, 10)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    cols = [('Pessoa', 22), ('Dia', 9), ('Mes', 9), ('Ano', 12), ('Descricao', 63),
            ('Categoria', 24), ('Pagamento', 28), ('Banco', 22), ('Valor (R$)', 24),
            ('Parcela', 18), ('Recorrente', 14)]
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_fill_color(13, 110, 253)
    pdf.set_text_color(255, 255, 255)
    for label, width in cols:
        pdf.cell(width, 7, label, border=1, fill=True, align='C')
    pdf.ln()
    pdf.set_font('Helvetica', size=8)
    pdf.set_text_color(0, 0, 0)
    for i, row in enumerate(rows):
        pdf.set_fill_color(*((240, 244, 255) if i % 2 == 0 else (255, 255, 255)))
        display = list(row)
        display[8] = brl(row[8])
        for j, (_, width) in enumerate(cols):
            pdf.cell(width, 7, _latin1(display[j])[:60], border=1, fill=True, align='R' if j == 8 else 'L')
        pdf.ln()
    return bytes(pdf.output())
