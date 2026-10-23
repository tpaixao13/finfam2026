from datetime import date, timedelta

import openpyxl
import io

from app.services.finance import month_offset


def _card(family, **kw):
    body = {'supports_credit': True, 'last_digits': '1234', 'bank': 'Nubank', 'due_day': 10,
            'best_buy_day': 3, 'credit_limit': '5000'}
    body.update(kw)
    r = family.post('/cards', json=body)
    assert r.status_code == 201, r.text
    return r.json()['card']


def _expense(family, **kw):
    body = {'description': 'Mercado', 'amount': '100.00', 'category': 'Alimentação',
            'payment_method': 'PIX', 'date': '2026-05-15'}
    body.update(kw)
    r = family.post('/expenses', json=body)
    assert r.status_code == 201, r.text
    return r.json()


def test_simple_expense_and_defaults(family):
    out = _expense(family)
    e = out['expenses'][0]
    assert out['kind'] == 'avista'
    assert (e['day'], e['month'], e['year']) == (15, 5, 2026)
    assert e['user_id'] == family.user['id']
    assert e['status'] in ('vencida', 'pendente')
    # sem data → hoje
    out = _expense(family, date=None)
    today = date.today()
    assert out['expenses'][0]['purchase_date'] == today.isoformat()


def test_validation_messages(family):
    r = family.post('/expenses', json={'description': 'x', 'amount': '0', 'category': 'A', 'payment_method': 'PIX'})
    assert r.status_code == 422
    assert r.json()['detail'].startswith('Valor')
    r = family.post('/expenses', json={'description': 'x', 'amount': '10', 'category': 'A',
                                       'payment_method': 'PIX', 'installments': 3})
    assert r.status_code == 422
    assert 'Crédito' in r.json()['detail']


def test_installments_split_and_billing_month(family):
    card = _card(family, best_buy_day=3, due_day=10)
    # compra dia 20/12 > melhor dia 3 → primeira parcela na fatura de janeiro
    out = _expense(family, description='TV', amount='100.00', payment_method='Cartão de Crédito',
                   card_id=card['id'], installments=3, date='2026-12-20')
    assert out['kind'] == 'parcelado'
    items = out['expenses']
    assert [i['amount'] for i in items] == [33.33, 33.33, 33.34]
    assert [(i['month'], i['year']) for i in items] == [(1, 2027), (2, 2027), (3, 2027)]
    assert items[0]['purchase_month'] == 12 and items[0]['purchase_year'] == 2026
    assert items[0]['series']['label'] == '1/3'
    assert items[0]['due_date'] == '2027-01-10'
    assert items[0]['bank'] == 'Nubank'


def test_credit_before_best_day_same_month(family):
    card = _card(family, best_buy_day=10, due_day=17)
    e = _expense(family, payment_method='Cartão de Crédito', card_id=card['id'], date='2026-05-08')['expenses'][0]
    assert (e['month'], e['year']) == (5, 2026)


def test_recurring_and_group_edit(family):
    out = _expense(family, description='Netflix', amount='55.90', category='Lazer', recurring_months=4)
    items = out['expenses']
    assert len(items) == 4 and items[-1]['month'] == 8
    gid = items[0]['series']['group_id']
    r = family.patch(f"/expenses/{items[0]['id']}", json={'amount': '59.90', 'category': 'Assinaturas',
                                                         'apply_to_group': True})
    assert r.status_code == 200 and 'todas as 4' in r.json()['message']
    lst = family.get('/expenses', params={'month': 0, 'year': 2026, 'category': 'Assinaturas'}).json()
    assert lst['total'] == 4 and all(i['amount'] == 59.9 for i in lst['items'])
    assert family.delete(f'/expenses/recurring-groups/{gid}').status_code == 200
    assert family.get('/expenses', params={'month': 0, 'year': 2026}).json()['total'] == 0


def test_duplicate_flag(family):
    _expense(family)
    out = _expense(family)
    assert out['duplicate'] is True


def test_patch_partial_and_move_date(family):
    e = _expense(family)['expenses'][0]
    r = family.patch(f"/expenses/{e['id']}", json={'day': 2})
    assert r.status_code == 200
    moved = r.json()['expense']
    assert (moved['day'], moved['month'], moved['description']) == (2, 5, 'Mercado')


def test_bills_and_pay(family):
    today = date.today()
    past = today - timedelta(days=3)
    soon = today + timedelta(days=2)
    late = _expense(family, description='Luz', date=past.isoformat())['expenses'][0]
    now = _expense(family, description='Água', date=today.isoformat())['expenses'][0]
    nxt = _expense(family, description='Internet', date=soon.isoformat())['expenses'][0]
    _expense(family, description='Pago', date=today.isoformat(), paid=True)
    far = today + timedelta(days=40)
    _expense(family, description='Longe', date=far.isoformat())

    b = family.get('/expenses/bills', params={'days_ahead': 7}).json()
    assert [e['id'] for e in b['overdue']] == [late['id']]
    assert [e['id'] for e in b['today']] == [now['id']]
    assert [e['id'] for e in b['upcoming']] == [nxt['id']]
    assert b['total'] == 300.0

    r = family.post('/expenses/pay', json={'ids': [late['id'], now['id']]})
    assert r.json()['changed'] == 2
    b = family.get('/expenses/bills').json()
    assert b['overdue'] == [] and b['today'] == []


def test_toggle_delete_and_installment_group_delete(family):
    e = _expense(family)['expenses'][0]
    assert family.post(f"/expenses/{e['id']}/toggle-paid").json()['expense']['paid'] is True
    assert family.delete(f"/expenses/{e['id']}").status_code == 200
    assert family.get(f"/expenses/{e['id']}").status_code == 404
    card = _card(family)
    out = _expense(family, payment_method='Cartão de Crédito', card_id=card['id'], installments=5)
    gid = out['expenses'][0]['series']['group_id']
    assert '5 parcelas' in family.delete(f'/expenses/installment-groups/{gid}').json()['message']


def test_filters_pagination_export(family):
    for i in range(25):
        _expense(family, description=f'Item {i}', amount=f'{i + 1}.00', date=f'2026-03-{(i % 28) + 1:02d}')
    _expense(family, description='Uber', category='Transporte', date='2026-03-05')
    r = family.get('/expenses', params={'month': 3, 'year': 2026, 'per_page': 10, 'page': 3}).json()
    assert r['total'] == 26 and r['pages'] == 3 and len(r['items']) == 6
    assert r['total_amount'] == sum(range(1, 26)) + 100
    r = family.get('/expenses', params={'month': 3, 'year': 2026, 'search': 'ube'}).json()
    assert r['total'] == 1
    r = family.get('/expenses/export', params={'month': 3, 'year': 2026, 'fmt': 'xlsx'})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 27 and ws['A1'].value == 'Pessoa'
    r = family.get('/expenses/export', params={'month': 3, 'year': 2026, 'fmt': 'pdf'})
    assert r.status_code == 200 and r.content.startswith(b'%PDF')
    opts = family.get('/expenses/options').json()
    assert 'Transporte' in opts['categories']


def test_suggest_category_rules_then_history(family):
    family.post('/rules', json={'keyword': 'ifood', 'category': 'Alimentação'})
    r = family.get('/expenses/suggest-category', params={'description': 'IFOOD *Restaurante'}).json()
    assert r == {'category': 'Alimentação', 'source': 'rule'}
    _expense(family, description='Posto Shell', category='Transporte')
    r = family.get('/expenses/suggest-category', params={'description': 'Posto Shell'}).json()
    assert r['category'] == 'Transporte'


def test_month_offset():
    assert month_offset(12, 2026, 1) == (1, 2027)
    assert month_offset(1, 2026, -1) == (12, 2025)
    assert month_offset(5, 2026, 14) == (7, 2027)
